from dataclasses import dataclass
from decimal import Decimal, InvalidOperation
from pathlib import Path
import re
import unicodedata

from django.conf import settings
from django.db import transaction
from django.utils import timezone
from openpyxl import load_workbook

from audit.models import EventoAuditoria
from audit.services import registrar_evento
from diagnostics.models import Diagnostico, TarefaProcessamento
from diagnostics.services import registrar_diagnostico
from files.models import Arquivo
from files.services import store_uploaded_file

from .models import ItemPedido, ModeloPlanilha, ModeloPlanilhaVersao, Pedido


class SpreadsheetImportError(Exception):
    def __init__(self, message: str, detail: str = ""):
        super().__init__(message)
        self.message = message
        self.detail = detail


@dataclass(frozen=True)
class ParsedOrderItem:
    sheet: str
    row_number: int
    code_raw: str
    code_normalized: str
    product_raw: str
    product_normalized: str
    quantity_raw: str
    quantity: Decimal
    unit_raw: str = ""
    unit_normalized: str = ""
    size_raw: str = ""
    size_normalized: str = ""
    estimated_price_raw: str = ""
    estimated_price: Decimal | None = None


@dataclass(frozen=True)
class OrderImportResult:
    arquivo: Arquivo
    pedido: Pedido | None
    created: bool
    duplicate: bool
    diagnostics_count: int = 0


HEADER_ALIASES = {
    "code": {"CODIGO DO ITEM", "CODIGO", "COD ITEM", "CODIGO ITEM", "ITEM", "COD"},
    "product": {"PRODUTO", "DESCRICAO", "DESCRICAO DO ITEM", "NOME DO PRODUTO", "MATERIAL"},
    "quantity": {"QUANTIDADE", "QTD", "QTDE", "QTD PEDIDA", "QUANT"},
    "unit": {"UNIDADE", "UN", "UND", "UNID"},
    "size": {"TAMANHO", "CONTEUDO", "MEDIDA", "EMBALAGEM"},
    "price": {"PRECO", "PRECO ESTIMADO", "VALOR", "VLR UNIT", "VALOR UNITARIO"},
}


def normalize_text(value) -> str:
    text = "" if value is None else str(value).strip()
    decomposed = unicodedata.normalize("NFKD", text)
    without_accents = "".join(char for char in decomposed if not unicodedata.combining(char))
    return re.sub(r"\s+", " ", without_accents).upper()


def normalize_code(value) -> str:
    text = normalize_text(value)
    cleaned = re.sub(r"[^A-Z0-9]+", "", text)
    if cleaned.isdigit():
        return cleaned.lstrip("0") or "0"
    return cleaned


def normalize_decimal(value) -> Decimal | None:
    if value is None:
        return None
    if isinstance(value, Decimal):
        return value
    if isinstance(value, int):
        return Decimal(value)
    if isinstance(value, float):
        return Decimal(str(value))

    text = str(value).strip()
    if not text:
        return None

    text = text.replace("R$", "").replace(" ", "")
    if "," in text and "." in text:
        text = text.replace(".", "").replace(",", ".")
    elif "," in text:
        text = text.replace(",", ".")

    try:
        return Decimal(text)
    except InvalidOperation as exc:
        raise SpreadsheetImportError("Valor numerico invalido na planilha.", f"Valor recebido: {value}") from exc


def original_value(value) -> str:
    if value is None:
        return ""
    return str(value).strip()


def find_header(row) -> dict[str, int]:
    normalized_cells = [normalize_text(cell.value) for cell in row]
    mapping: dict[str, int] = {}
    for field, aliases in HEADER_ALIASES.items():
        for index, value in enumerate(normalized_cells):
            if value in aliases:
                mapping[field] = index
                break
    return mapping


def detect_header(sheet) -> tuple[int, dict[str, int]]:
    for row in sheet.iter_rows(min_row=1, max_row=30):
        mapping = find_header(row)
        if {"code", "product", "quantity"}.issubset(mapping):
            return row[0].row, mapping
    raise SpreadsheetImportError(
        "Estrutura da planilha nao reconhecida.",
        "Cabecalhos obrigatorios nao encontrados: CODIGO DO ITEM, PRODUTO/DESCRICAO e QUANTIDADE.",
    )


def detect_order_number(workbook) -> str:
    patterns = [re.compile(r"\bW\d{6,}\b", re.IGNORECASE), re.compile(r"\bP-\d{4}-\d+\b", re.IGNORECASE)]
    for sheet in workbook.worksheets:
        for row in sheet.iter_rows(min_row=1, max_row=min(sheet.max_row, 20)):
            for cell in row:
                value = original_value(cell.value)
                for pattern in patterns:
                    match = pattern.search(value)
                    if match:
                        return match.group(0).upper()
    return ""


def parse_coral_workbook(path: Path) -> tuple[str, list[ParsedOrderItem]]:
    try:
        workbook = load_workbook(filename=path, data_only=True, read_only=True)
    except Exception as exc:
        raise SpreadsheetImportError("Planilha ilegivel.", str(exc)) from exc

    try:
        items: list[ParsedOrderItem] = []
        negative_rows: list[str] = []
        for sheet in workbook.worksheets:
            header_row, mapping = detect_header(sheet)
            for row in sheet.iter_rows(min_row=header_row + 1):
                quantity_raw = original_value(row[mapping["quantity"]].value)
                quantity = normalize_decimal(row[mapping["quantity"]].value)
                if quantity is None or quantity == 0:
                    continue
                if quantity < 0:
                    negative_rows.append(f"{sheet.title}!{row[0].row}")
                    continue

                code_raw = original_value(row[mapping["code"]].value)
                product_raw = original_value(row[mapping["product"]].value)
                if not code_raw and not product_raw:
                    continue

                price_raw = original_value(row[mapping["price"]].value) if "price" in mapping else ""
                estimated_price = normalize_decimal(row[mapping["price"]].value) if "price" in mapping else None
                unit_raw = original_value(row[mapping["unit"]].value) if "unit" in mapping else ""
                size_raw = original_value(row[mapping["size"]].value) if "size" in mapping else ""

                items.append(
                    ParsedOrderItem(
                        sheet=sheet.title,
                        row_number=row[0].row,
                        code_raw=code_raw,
                        code_normalized=normalize_code(code_raw),
                        product_raw=product_raw,
                        product_normalized=normalize_text(product_raw),
                        quantity_raw=quantity_raw,
                        quantity=quantity,
                        unit_raw=unit_raw,
                        unit_normalized=normalize_text(unit_raw),
                        size_raw=size_raw,
                        size_normalized=normalize_text(size_raw),
                        estimated_price_raw=price_raw,
                        estimated_price=estimated_price,
                    )
                )

        if negative_rows:
            raise SpreadsheetImportError(
                "Planilha possui quantidade negativa e precisa de revisao.",
                f"Linhas com quantidade negativa: {', '.join(negative_rows)}",
            )
        if not items:
            raise SpreadsheetImportError(
                "Planilha sem itens validos.",
                "Nenhuma linha com quantidade positiva foi encontrada.",
            )

        return detect_order_number(workbook), items
    finally:
        workbook.close()


def get_or_create_coral_model(*, empresa_cliente, fornecedor) -> ModeloPlanilhaVersao:
    modelo, _ = ModeloPlanilha.objects.get_or_create(
        empresa_cliente=empresa_cliente,
        fornecedor=fornecedor,
        nome="Coral",
    )
    versao, _ = ModeloPlanilhaVersao.objects.get_or_create(
        modelo=modelo,
        versao=1,
        defaults={"configuracao": {"adapter": "coral_v1"}},
    )
    return versao


def import_coral_order_spreadsheet(*, empresa_cliente, loja, fornecedor, uploaded_file, usuario) -> OrderImportResult:
    stored = store_uploaded_file(
        empresa_cliente=empresa_cliente,
        uploaded_file=uploaded_file,
        usuario=usuario,
        tipo=Arquivo.Tipo.PLANILHA_PEDIDO,
    )

    if stored.duplicate:
        pedido = Pedido.objects.filter(arquivo=stored.arquivo).first()
        return OrderImportResult(arquivo=stored.arquivo, pedido=pedido, created=False, duplicate=True)

    tarefa = TarefaProcessamento.objects.create(
        empresa_cliente=empresa_cliente,
        arquivo=stored.arquivo,
        tipo=TarefaProcessamento.Tipo.IMPORTAR_PLANILHA,
        estado=TarefaProcessamento.Estado.PROCESSANDO,
        progresso=10,
    )
    path = settings.NFCH_STORAGE_ROOT / stored.arquivo.caminho_relativo

    try:
        numero_pedido, parsed_items = parse_coral_workbook(path)
        with transaction.atomic():
            modelo_versao = get_or_create_coral_model(empresa_cliente=empresa_cliente, fornecedor=fornecedor)
            pedido = Pedido.objects.create(
                empresa_cliente=empresa_cliente,
                loja=loja,
                fornecedor=fornecedor,
                arquivo=stored.arquivo,
                modelo_versao=modelo_versao,
                numero_pedido=numero_pedido,
                data_operacional=timezone.localdate(),
            )
            ItemPedido.objects.bulk_create(
                [
                    ItemPedido(
                        pedido=pedido,
                        empresa_cliente=empresa_cliente,
                        fornecedor=fornecedor,
                        aba=item.sheet,
                        linha=item.row_number,
                        codigo_original=item.code_raw,
                        codigo_normalizado=item.code_normalized,
                        produto_original=item.product_raw,
                        produto_normalizado=item.product_normalized,
                        quantidade_original=item.quantity_raw,
                        quantidade_pedida=item.quantity,
                        unidade_original=item.unit_raw,
                        unidade_normalizada=item.unit_normalized,
                        tamanho_original=item.size_raw,
                        tamanho_normalizado=item.size_normalized,
                        preco_estimado_original=item.estimated_price_raw,
                        preco_estimado=item.estimated_price,
                        saldo_cache=item.quantity,
                    )
                    for item in parsed_items
                ]
            )
            tarefa.estado = TarefaProcessamento.Estado.CONCLUIDO
            tarefa.progresso = 100
            tarefa.mensagem_usuario = f"Pedido importado com {len(parsed_items)} item(ns)."
            tarefa.save(update_fields=["estado", "progresso", "mensagem_usuario", "atualizada_em"])
            registrar_evento(
                empresa_cliente=empresa_cliente,
                usuario=usuario,
                origem=EventoAuditoria.Origem.USUARIO,
                entidade=pedido,
                acao="PEDIDO_IMPORTADO",
                depois={
                    "arquivo_id": stored.arquivo.id,
                    "numero_pedido": pedido.numero_pedido,
                    "itens": len(parsed_items),
                },
            )
        return OrderImportResult(arquivo=stored.arquivo, pedido=pedido, created=True, duplicate=False)
    except SpreadsheetImportError as exc:
        stored.arquivo.situacao = Arquivo.Situacao.ERRO
        stored.arquivo.save(update_fields=["situacao"])
        tarefa.estado = TarefaProcessamento.Estado.ERRO
        tarefa.progresso = 100
        tarefa.mensagem_usuario = exc.message
        tarefa.erro_tecnico = exc.detail
        tarefa.save(update_fields=["estado", "progresso", "mensagem_usuario", "erro_tecnico", "atualizada_em"])
        registrar_diagnostico(
            empresa_cliente=empresa_cliente,
            arquivo=stored.arquivo,
            tipo=Diagnostico.Tipo.LINHA_INVALIDA if "negativa" in exc.message else Diagnostico.Tipo.ESTRUTURA_NAO_RECONHECIDA,
            mensagem_usuario=exc.message,
            detalhe_tecnico=exc.detail,
            criado_por=usuario,
        )
        return OrderImportResult(arquivo=stored.arquivo, pedido=None, created=False, duplicate=False, diagnostics_count=1)
