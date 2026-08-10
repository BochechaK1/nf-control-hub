from io import BytesIO
from decimal import Decimal
from pathlib import Path
from tempfile import TemporaryDirectory

from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import TestCase, override_settings
from django.urls import reverse
from django.utils import timezone
from openpyxl import Workbook

from accounts.models import Usuario
from approvals.models import Associacao
from files.models import Arquivo
from invoices.models import ItemNotaFiscal, NotaFiscal
from matching.models import Conferencia, ConferenciaCandidata, ConferenciaItem
from orders.models import ItemPedido, ModeloPlanilha, ModeloPlanilhaVersao, Pedido
from organizations.models import EmpresaCliente, Fornecedor, Loja
from reports.models import Exportacao


class UiPreviewTests(TestCase):
    def setUp(self):
        self.user = Usuario.objects.create_superuser(username="kauan", password="secret")
        self.client.force_login(self.user)

    def test_dashboard_renders_operational_preview(self):
        response = self.client.get(reverse("ui:dashboard"))

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "NF Control Hub")
        self.assertContains(response, "Fila prioritaria")
        self.assertContains(response, "activity-disclosure")
        self.assertContains(response, "Ver historico")
        self.assertContains(response, 'aria-current="page"')
        self.assertContains(response, "Indicadores e fila")
        self.assertContains(response, "Nenhuma NF aguardando decisao.")
        self.assertNotContains(response, "Eventos auditaveis mais novos.")
        self.assertNotContains(response, "1.987.853")

    def test_custom_login_renders(self):
        self.client.logout()

        response = self.client.get(reverse("ui:login"))

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "Entrar no sistema")
        self.assertContains(response, "login-form")

    def test_mvp_menu_pages_render(self):
        names = [
            "ui:fila",
            "ui:recebimento",
            "ui:planilhas",
            "ui:relatorios",
            "ui:diagnostico",
            "ui:configuracoes",
            "ui:sobre",
        ]

        for name in names:
            with self.subTest(name=name):
                response = self.client.get(reverse(name))
                self.assertEqual(response.status_code, 200)

    def test_upload_forms_use_custom_file_picker(self):
        planilhas = self.client.get(reverse("ui:planilhas"))
        fila = self.client.get(reverse("ui:fila"))

        self.assertContains(planilhas, "file-picker")
        self.assertContains(planilhas, "Selecionar arquivos")
        self.assertContains(planilhas, "Importar planilhas de pedido")
        self.assertContains(planilhas, "Identificar automaticamente")
        self.assertContains(planilhas, "Pedidos ativos")
        self.assertContains(planilhas, "Saldo aberto")
        self.assertContains(planilhas, "Pedidos importados")
        self.assertNotContains(planilhas, 'name="empresa_cliente"')
        self.assertNotContains(planilhas, "Empresa opcional")
        self.assertNotContains(planilhas, 'name="loja" required')
        self.assertNotContains(planilhas, 'name="fornecedor" required')
        self.assertNotContains(planilhas, "Importar planilhas Coral")
        self.assertNotContains(planilhas, "Excel Coral")
        self.assertContains(fila, "file-picker")
        self.assertContains(fila, "Selecionar XML")
        self.assertNotContains(fila, 'name="empresa_cliente"')
        self.assertNotContains(fila, "Empresa</span>")
        self.assertContains(fila, "queue-summary")
        self.assertContains(fila, "Total na consulta")
        self.assertContains(fila, "Notas fiscais")
        self.assertContains(fila, "queue-list-panel")
        self.assertNotContains(fila, "queue-detail-panel")
        self.assertNotContains(fila, "Nenhuma NF selecionada")

    def test_recebimento_renders_operational_summary(self):
        response = self.client.get(reverse("ui:recebimento"))

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "receiving-summary")
        self.assertContains(response, "Confirmacoes pendentes")
        self.assertContains(response, "Historico de recebimento")
        self.assertContains(response, "Alerta 30 dias")

    def test_relatorios_renders_operational_sections(self):
        response = self.client.get(reverse("ui:relatorios"))

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "reports-summary")
        self.assertContains(response, "Itens faturados")
        self.assertContains(response, "Alocacoes aprovadas vigentes")
        self.assertContains(response, "Exportacoes")
        self.assertContains(response, "Historico recente")
        self.assertNotContains(response, ">A5<")

    def test_diagnostico_renders_triage_sections(self):
        response = self.client.get(reverse("ui:diagnostico"))

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "diagnostic-summary")
        self.assertContains(response, "Ocorrencias tecnicas")
        self.assertContains(response, "Tarefas")
        self.assertContains(response, "Fila e Diagnostico nao sao a mesma coisa")
        self.assertNotContains(response, "Reprocessar</button>")

    def test_visualizador_cannot_access_diagnostico(self):
        empresa = EmpresaCliente.objects.create(nome="Cliente", codigo="cliente")
        viewer = Usuario.objects.create_user(
            username="viewer",
            password="secret",
            role=Usuario.Role.VISUALIZADOR,
            empresa_cliente=empresa,
        )
        self.client.force_login(viewer)

        response = self.client.get(reverse("ui:diagnostico"))

        self.assertEqual(response.status_code, 403)

    def test_configuracoes_renders_real_registry_sections(self):
        response = self.client.get(reverse("ui:configuracoes"))

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "settings-summary")
        self.assertContains(response, "Empresas cliente")
        self.assertContains(response, "Lojas")
        self.assertContains(response, "Fornecedores")
        self.assertContains(response, "Usuarios")
        self.assertContains(response, "Configuracoes sao restritas")
        self.assertNotContains(response, "Nova empresa</button>")
        self.assertNotContains(response, "Alias</button>")

    def test_visualizador_cannot_access_configuracoes(self):
        empresa = EmpresaCliente.objects.create(nome="Cliente", codigo="cliente-config")
        viewer = Usuario.objects.create_user(
            username="viewer-config",
            password="secret",
            role=Usuario.Role.VISUALIZADOR,
            empresa_cliente=empresa,
        )
        self.client.force_login(viewer)

        response = self.client.get(reverse("ui:configuracoes"))

        self.assertEqual(response.status_code, 403)

    def test_planilhas_auto_identifies_context_from_filename(self):
        empresa = EmpresaCliente.objects.create(nome="Destiny In", codigo="destiny")
        loja = Loja.objects.create(empresa_cliente=empresa, codigo="matriz", nome="Matriz")
        fornecedor = Fornecedor.objects.create(
            empresa_cliente=empresa,
            nome="Coral",
            cnpj="11.111.111/0001-11",
        )
        workbook = Workbook()
        sheet = workbook.active
        sheet.append(["CODIGO DO ITEM", "PRODUTO", "QUANTIDADE", "UNIDADE"])
        sheet.append(["000123", "Produto teste", 4, "UN"])
        payload = BytesIO()
        workbook.save(payload)
        payload.seek(0)
        upload = SimpleUploadedFile(
            "CORAL MATRIZ (1).xlsx",
            payload.read(),
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )

        with TemporaryDirectory() as temp_dir, override_settings(NFCH_STORAGE_ROOT=Path(temp_dir)):
            response = self.client.post(reverse("ui:planilhas"), {"planilhas": upload}, follow=True)

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "pedido importado com sucesso")
        self.assertTrue(
            Pedido.objects.filter(
                empresa_cliente=empresa,
                loja=loja,
                fornecedor=fornecedor,
                numero_pedido="",
            ).exists()
        )

    def test_planilhas_do_not_assign_unknown_supplier_to_only_registered_supplier(self):
        empresa = EmpresaCliente.objects.create(nome="Destiny In", codigo="destiny-unknown-supplier")
        Loja.objects.create(empresa_cliente=empresa, codigo="paraty", nome="paraty")
        Fornecedor.objects.create(
            empresa_cliente=empresa,
            nome="Akzo Nobel Ltda / Coral",
            cnpj="60.561.719/0095-03",
        )
        workbook = Workbook()
        sheet = workbook.active
        sheet.append([None, "Produto", "Descrição", None, None, None, None, None, None, None, "QUANTID."])
        sheet.append([None, "613305730R", "CORANTE LIQUIDO PRETO 50ML", None, None, None, None, None, None, None, 36])
        payload = BytesIO()
        workbook.save(payload)
        payload.seek(0)
        upload = SimpleUploadedFile(
            "IQUINE PARATY (4).xlsx",
            payload.read(),
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )

        with TemporaryDirectory() as temp_dir, override_settings(NFCH_STORAGE_ROOT=Path(temp_dir)):
            response = self.client.post(reverse("ui:planilhas"), {"planilhas": upload}, follow=True)

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "fornecedor nao cadastrado ou nao identificado")
        self.assertFalse(Pedido.objects.filter(empresa_cliente=empresa).exists())

    def test_planilhas_identifies_supplier_by_distinctive_token(self):
        empresa = EmpresaCliente.objects.create(nome="Destiny In", codigo="destiny-iquine-token")
        loja = Loja.objects.create(empresa_cliente=empresa, codigo="paraty", nome="paraty")
        fornecedor = Fornecedor.objects.create(
            empresa_cliente=empresa,
            nome="TINTAS IQUINE LTDA",
            cnpj="09.722.463/0006-46",
        )
        workbook = Workbook()
        sheet = workbook.active
        sheet.append([None, "Produto", "Descricao", None, None, None, None, None, None, None, "QUANTID."])
        sheet.append([None, "613305730R", "CORANTE LIQUIDO PRETO 50ML", None, None, None, None, None, None, None, 36])
        payload = BytesIO()
        workbook.save(payload)
        payload.seek(0)
        upload = SimpleUploadedFile(
            "IQUINE PARATY (4).xlsx",
            payload.read(),
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )

        with TemporaryDirectory() as temp_dir, override_settings(NFCH_STORAGE_ROOT=Path(temp_dir)):
            response = self.client.post(reverse("ui:planilhas"), {"planilhas": upload}, follow=True)

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "pedido importado com sucesso")
        self.assertTrue(Pedido.objects.filter(empresa_cliente=empresa, loja=loja, fornecedor=fornecedor).exists())

    def test_xml_import_recovers_stored_iquine_spreadsheet_and_recalculates_match(self):
        empresa = EmpresaCliente.objects.create(nome="Destiny In", codigo="destiny-iquine-recover")
        loja = Loja.objects.create(
            empresa_cliente=empresa,
            codigo="paraty",
            nome="paraty",
            cnpj="09.580.958/0002-54",
        )
        workbook = Workbook()
        sheet = workbook.active
        sheet.append([None, "Produto", "Descricao", None, None, None, None, None, None, None, "QUANTID."])
        sheet.append([None, "613305730R", "CORANTE LIQUIDO PRETO 50ML", None, None, None, None, None, None, None, 36])
        payload = BytesIO()
        workbook.save(payload)
        stored_bytes = payload.getvalue()
        xml = b"""<?xml version="1.0" encoding="UTF-8"?>
<nfeProc xmlns="http://www.portalfiscal.inf.br/nfe" versao="4.00">
  <NFe>
    <infNFe Id="NFe33260809722463000646550010000434401938843899" versao="4.00">
      <ide><natOp>Venda</natOp><mod>55</mod><serie>1</serie><nNF>43440</nNF><dhEmi>2026-08-06T10:30:00-03:00</dhEmi></ide>
      <emit><CNPJ>09722463000646</CNPJ><xNome>TINTAS IQUINE LTDA</xNome></emit>
      <dest><CNPJ>09580958000254</CNPJ><xNome>paraty</xNome></dest>
      <det nItem="1"><prod><cProd>613305730R</cProd><cEAN>SEM GTIN</cEAN><xProd>CORANTE LIQUIDO PRETO 50ML</xProd><NCM>32091010</NCM><CFOP>5102</CFOP><uCom>UN</uCom><qCom>36.0000</qCom><vUnCom>2.020000</vUnCom><vProd>72.72</vProd></prod></det>
      <total><ICMSTot><vNF>72.72</vNF></ICMSTot></total>
    </infNFe>
  </NFe>
  <protNFe versao="4.00"><infProt><chNFe>33260809722463000646550010000434401938843899</chNFe><cStat>100</cStat></infProt></protNFe>
</nfeProc>
"""
        with TemporaryDirectory() as temp_dir, override_settings(NFCH_STORAGE_ROOT=Path(temp_dir)):
            relative_path = Path("planilha_pedido/2026/08/iquine_paraty.xlsx")
            absolute_path = Path(temp_dir) / relative_path
            absolute_path.parent.mkdir(parents=True, exist_ok=True)
            absolute_path.write_bytes(stored_bytes)
            Arquivo.objects.create(
                empresa_cliente=empresa,
                tipo=Arquivo.Tipo.PLANILHA_PEDIDO,
                nome_original="IQUINE PARATY (4).xlsx",
                tamanho=len(stored_bytes),
                hash_sha256="iquine-paraty-stored".ljust(64, "0"),
                caminho_relativo=relative_path.as_posix(),
                usuario_importacao=self.user,
            )
            upload = SimpleUploadedFile("iquine.xml", xml, content_type="application/xml")

            response = self.client.post(reverse("ui:fila"), {"xmls": upload}, follow=True)

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "planilha(s) armazenada(s) reprocessada(s)")
        pedido = Pedido.objects.get(empresa_cliente=empresa)
        nota = NotaFiscal.objects.get(empresa_cliente=empresa)
        self.assertEqual(pedido.fornecedor.nome, "TINTAS IQUINE LTDA")
        self.assertEqual(pedido.loja, loja)
        candidata = ConferenciaCandidata.objects.get(conferencia__nota_fiscal=nota)
        self.assertEqual(candidata.compatibilidade, Decimal("100.00"))

    def test_xml_import_recovers_stored_spreadsheet_by_item_codes_for_any_supplier(self):
        empresa = EmpresaCliente.objects.create(nome="Destiny In", codigo="destiny-any-supplier")
        loja = Loja.objects.create(
            empresa_cliente=empresa,
            codigo="paraty",
            nome="paraty",
            cnpj="09.580.958/0002-54",
        )
        workbook = Workbook()
        sheet = workbook.active
        sheet.append(["Codigo", "Descricao", "Quantidade"])
        sheet.append(["ALF100", "Produto Alfa 100", 10])
        sheet.append(["ALF200", "Produto Alfa 200", 5])
        payload = BytesIO()
        workbook.save(payload)
        stored_bytes = payload.getvalue()
        key = "9" * 44
        xml = f"""<?xml version="1.0" encoding="UTF-8"?>
<nfeProc xmlns="http://www.portalfiscal.inf.br/nfe" versao="4.00">
  <NFe>
    <infNFe Id="NFe{key}" versao="4.00">
      <ide><natOp>Venda</natOp><mod>55</mod><serie>1</serie><nNF>90001</nNF><dhEmi>2026-08-06T10:30:00-03:00</dhEmi></ide>
      <emit><CNPJ>12345678000199</CNPJ><xNome>FABRICA ALFA LTDA</xNome></emit>
      <dest><CNPJ>09580958000254</CNPJ><xNome>paraty</xNome></dest>
      <det nItem="1"><prod><cProd>ALF100</cProd><cEAN>SEM GTIN</cEAN><xProd>Produto Alfa 100</xProd><NCM>32091010</NCM><CFOP>5102</CFOP><uCom>UN</uCom><qCom>10.0000</qCom><vUnCom>1.000000</vUnCom><vProd>10.00</vProd></prod></det>
      <det nItem="2"><prod><cProd>ALF200</cProd><cEAN>SEM GTIN</cEAN><xProd>Produto Alfa 200</xProd><NCM>32091010</NCM><CFOP>5102</CFOP><uCom>UN</uCom><qCom>5.0000</qCom><vUnCom>1.000000</vUnCom><vProd>5.00</vProd></prod></det>
      <total><ICMSTot><vNF>15.00</vNF></ICMSTot></total>
    </infNFe>
  </NFe>
  <protNFe versao="4.00"><infProt><chNFe>{key}</chNFe><cStat>100</cStat></infProt></protNFe>
</nfeProc>
""".encode("utf-8")
        with TemporaryDirectory() as temp_dir, override_settings(NFCH_STORAGE_ROOT=Path(temp_dir)):
            relative_path = Path("planilha_pedido/2026/08/paraty_alfa.xlsx")
            absolute_path = Path(temp_dir) / relative_path
            absolute_path.parent.mkdir(parents=True, exist_ok=True)
            absolute_path.write_bytes(stored_bytes)
            Arquivo.objects.create(
                empresa_cliente=empresa,
                tipo=Arquivo.Tipo.PLANILHA_PEDIDO,
                nome_original="PEDIDO PARATY 001.xlsx",
                tamanho=len(stored_bytes),
                hash_sha256="any-supplier-stored".ljust(64, "0"),
                caminho_relativo=relative_path.as_posix(),
                usuario_importacao=self.user,
            )
            upload = SimpleUploadedFile("alfa.xml", xml, content_type="application/xml")

            response = self.client.post(reverse("ui:fila"), {"xmls": upload}, follow=True)

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "planilha(s) armazenada(s) reprocessada(s)")
        fornecedor = Fornecedor.objects.get(empresa_cliente=empresa)
        self.assertEqual(fornecedor.nome, "FABRICA ALFA LTDA")
        pedido = Pedido.objects.get(empresa_cliente=empresa)
        self.assertEqual(pedido.fornecedor, fornecedor)
        candidata = ConferenciaCandidata.objects.get(conferencia__nota_fiscal__empresa_cliente=empresa)
        self.assertEqual(candidata.compatibilidade, Decimal("100.00"))

    def test_planilhas_summary_counts_distinct_orders(self):
        empresa = EmpresaCliente.objects.create(nome="Cliente", codigo="cliente-summary")
        loja = Loja.objects.create(empresa_cliente=empresa, codigo="matriz", nome="Matriz")
        fornecedor = Fornecedor.objects.create(
            empresa_cliente=empresa,
            nome="Coral",
            cnpj="22.222.222/0001-22",
        )
        modelo = ModeloPlanilha.objects.create(empresa_cliente=empresa, fornecedor=fornecedor, nome="Coral")
        versao = ModeloPlanilhaVersao.objects.create(modelo=modelo, versao=1)
        pedidos = []
        for index in range(2):
            arquivo = Arquivo.objects.create(
                empresa_cliente=empresa,
                tipo=Arquivo.Tipo.PLANILHA_PEDIDO,
                nome_original=f"pedido-{index}.xlsx",
                tamanho=10,
                hash_sha256=f"hash-{index}",
                caminho_relativo=f"planilhas/pedido-{index}.xlsx",
                usuario_importacao=self.user,
            )
            pedidos.append(
                Pedido.objects.create(
                    empresa_cliente=empresa,
                    loja=loja,
                    fornecedor=fornecedor,
                    arquivo=arquivo,
                    modelo_versao=versao,
                    numero_pedido=f"P-{index}",
                    data_operacional=timezone.localdate(),
                )
            )
        ItemPedido.objects.create(
            pedido=pedidos[0],
            empresa_cliente=empresa,
            fornecedor=fornecedor,
            aba="Pedido",
            linha=1,
            codigo_original="1",
            codigo_normalizado="1",
            produto_original="Produto 1",
            produto_normalizado="PRODUTO 1",
            quantidade_original="2",
            quantidade_pedida=Decimal("2.000"),
            saldo_cache=Decimal("2.000"),
        )
        ItemPedido.objects.create(
            pedido=pedidos[0],
            empresa_cliente=empresa,
            fornecedor=fornecedor,
            aba="Pedido",
            linha=2,
            codigo_original="2",
            codigo_normalizado="2",
            produto_original="Produto 2",
            produto_normalizado="PRODUTO 2",
            quantidade_original="3",
            quantidade_pedida=Decimal("3.000"),
            saldo_cache=Decimal("3.000"),
        )
        ItemPedido.objects.create(
            pedido=pedidos[1],
            empresa_cliente=empresa,
            fornecedor=fornecedor,
            aba="Pedido",
            linha=1,
            codigo_original="3",
            codigo_normalizado="3",
            produto_original="Produto 3",
            produto_normalizado="PRODUTO 3",
            quantidade_original="4",
            quantidade_pedida=Decimal("4.000"),
            saldo_cache=Decimal("4.000"),
        )

        response = self.client.get(reverse("ui:planilhas"))

        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.context["resumo"]["pedidos_total"], 2)
        self.assertEqual(response.context["resumo"]["itens_total"], 3)
        self.assertContains(response, "<strong>9</strong>", html=True)
        self.assertContains(response, "Recebido")
        self.assertNotContains(response, "9,000")
        self.assertNotContains(response, "0,000")

    def test_fila_renders_extra_and_absent_comparison_rows(self):
        empresa = EmpresaCliente.objects.create(nome="Cliente", codigo="cliente-fila")
        loja = Loja.objects.create(empresa_cliente=empresa, codigo="matriz", nome="Matriz", cnpj="11.111.111/0001-11")
        fornecedor = Fornecedor.objects.create(
            empresa_cliente=empresa,
            nome="Akzo Nobel Ltda / Coral",
            cnpj="60.561.719/0095-03",
        )
        modelo = ModeloPlanilha.objects.create(empresa_cliente=empresa, fornecedor=fornecedor, nome="Coral")
        versao = ModeloPlanilhaVersao.objects.create(modelo=modelo, versao=1)
        arquivo_pedido = Arquivo.objects.create(
            empresa_cliente=empresa,
            tipo=Arquivo.Tipo.PLANILHA_PEDIDO,
            nome_original="CORAL MATRIZ.xlsx",
            tamanho=10,
            hash_sha256="hash-planilha-fila",
            caminho_relativo="planilhas/coral.xlsx",
            usuario_importacao=self.user,
        )
        pedido = Pedido.objects.create(
            empresa_cliente=empresa,
            loja=loja,
            fornecedor=fornecedor,
            arquivo=arquivo_pedido,
            modelo_versao=versao,
            data_operacional=timezone.localdate(),
        )
        item_pedido = ItemPedido.objects.create(
            pedido=pedido,
            empresa_cliente=empresa,
            fornecedor=fornecedor,
            aba="Pedido",
            linha=1,
            codigo_original="5200001",
            codigo_normalizado="5200001",
            produto_original="Produto ausente na NF",
            produto_normalizado="PRODUTO AUSENTE NA NF",
            quantidade_original="3",
            quantidade_pedida=Decimal("3.000"),
            saldo_cache=Decimal("3.000"),
        )
        arquivo_xml = Arquivo.objects.create(
            empresa_cliente=empresa,
            tipo=Arquivo.Tipo.XML_NFE,
            nome_original="nfe.xml",
            tamanho=10,
            hash_sha256="hash-xml-fila",
            caminho_relativo="xml/nfe.xml",
            usuario_importacao=self.user,
        )
        nota = NotaFiscal.objects.create(
            empresa_cliente=empresa,
            arquivo_xml=arquivo_xml,
            fornecedor=fornecedor,
            loja=loja,
            chave_acesso="1" * 44,
            modelo="55",
            numero="1986570",
            serie="21",
            data_emissao=timezone.now(),
            natureza_operacao="Venda",
            emitente_cnpj=fornecedor.cnpj_normalizado,
            emitente_nome=fornecedor.nome,
            destinatario_cnpj=loja.cnpj_normalizado,
            destinatario_nome=loja.nome,
            valor_total=Decimal("10.00"),
        )
        item_nf = ItemNotaFiscal.objects.create(
            nota_fiscal=nota,
            empresa_cliente=empresa,
            numero_item=1,
            codigo_original="999999",
            codigo_normalizado="999999",
            descricao_original="Produto extra na NF",
            descricao_normalizada="PRODUTO EXTRA NA NF",
            ean_original="SEM GTIN",
            ncm="32091010",
            cfop="5102",
            unidade_original="UN",
            unidade_normalizada="UN",
            quantidade_original="1.0000",
            quantidade=Decimal("1.0000"),
        )
        conferencia = Conferencia.objects.create(empresa_cliente=empresa, nota_fiscal=nota)
        candidata = ConferenciaCandidata.objects.create(
            conferencia=conferencia,
            pedido=pedido,
            cobertura_itens=Decimal("0.00"),
            cobertura_quantidades=Decimal("0.00"),
            compatibilidade=Decimal("0.00"),
            nivel=ConferenciaCandidata.Nivel.BAIXA,
            alertas=["EXTRA_NF"],
        )
        ConferenciaItem.objects.create(
            candidata=candidata,
            item_nota_fiscal=item_nf,
            resultado=ConferenciaItem.Resultado.EXTRA_NF,
            quantidade_nf=Decimal("1.0000"),
            observacao="Item da NF ausente no pedido.",
            ordem=1,
        )
        ConferenciaItem.objects.create(
            candidata=candidata,
            item_pedido=item_pedido,
            resultado=ConferenciaItem.Resultado.SALDO_RESTANTE,
            saldo_pedido=Decimal("3.0000"),
            observacao="Item do pedido permanece em saldo para faturamento futuro.",
            ordem=2,
        )

        response = self.client.get(reverse("ui:fila"))

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "queue-card-disclosure")
        self.assertContains(response, "queue-top-actions")
        self.assertContains(response, "queue-action-approve")
        self.assertContains(response, "queue-action-reject")
        self.assertContains(response, 'data-toggle-details aria-expanded="false"')
        self.assertNotContains(response, "queue-detail-panel")
        self.assertContains(response, "Produto extra na NF")
        self.assertContains(response, "Produto ausente na NF")
        self.assertContains(response, "Aprovar")
        self.assertContains(response, "Reprovar")
        self.assertContains(response, "Extra na NF")
        self.assertContains(response, "Saldo restante")
        self.assertNotContains(response, "AUSENTE_NF")

    def test_fila_get_does_not_generate_missing_export(self):
        empresa = EmpresaCliente.objects.create(nome="Cliente", codigo="cliente-fila-no-write")
        loja = Loja.objects.create(empresa_cliente=empresa, codigo="matriz", nome="Matriz", cnpj="11.111.111/0001-11")
        fornecedor = Fornecedor.objects.create(
            empresa_cliente=empresa,
            nome="Fornecedor",
            cnpj="22.222.222/0001-22",
        )
        arquivo_pedido = Arquivo.objects.create(
            empresa_cliente=empresa,
            tipo=Arquivo.Tipo.PLANILHA_PEDIDO,
            nome_original="PEDIDO MATRIZ.xlsx",
            tamanho=10,
            hash_sha256="hash-fila-no-write-pedido",
            caminho_relativo="planilhas/pedido.xlsx",
            usuario_importacao=self.user,
        )
        pedido = Pedido.objects.create(
            empresa_cliente=empresa,
            loja=loja,
            fornecedor=fornecedor,
            arquivo=arquivo_pedido,
            data_operacional=timezone.localdate(),
        )
        arquivo_xml = Arquivo.objects.create(
            empresa_cliente=empresa,
            tipo=Arquivo.Tipo.XML_NFE,
            nome_original="nfe.xml",
            tamanho=10,
            hash_sha256="hash-fila-no-write-xml",
            caminho_relativo="xml/nfe.xml",
            usuario_importacao=self.user,
        )
        nota = NotaFiscal.objects.create(
            empresa_cliente=empresa,
            arquivo_xml=arquivo_xml,
            fornecedor=fornecedor,
            loja=loja,
            chave_acesso="2" * 44,
            modelo="55",
            numero="200",
            serie="1",
            data_emissao=timezone.now(),
            emitente_cnpj=fornecedor.cnpj_normalizado,
            emitente_nome=fornecedor.nome,
            destinatario_cnpj=loja.cnpj_normalizado,
            destinatario_nome=loja.nome,
            status_conferencia=NotaFiscal.StatusConferencia.APROVADA,
        )
        conferencia = Conferencia.objects.create(empresa_cliente=empresa, nota_fiscal=nota)
        candidata = ConferenciaCandidata.objects.create(
            conferencia=conferencia,
            pedido=pedido,
            cobertura_itens=Decimal("100.00"),
            cobertura_quantidades=Decimal("100.00"),
            compatibilidade=Decimal("100.00"),
            nivel=ConferenciaCandidata.Nivel.ALTA,
        )
        Associacao.objects.create(
            empresa_cliente=empresa,
            nota_fiscal=nota,
            pedido=pedido,
            candidata=candidata,
            aprovada_por=self.user,
        )

        response = self.client.get(reverse("ui:fila"))

        self.assertEqual(response.status_code, 200)
        self.assertEqual(Exportacao.objects.filter(empresa_cliente=empresa).count(), 0)

    def test_visualizador_menu_hides_admin_only_entries(self):
        empresa = EmpresaCliente.objects.create(nome="Cliente", codigo="cliente-menu")
        viewer = Usuario.objects.create_user(
            username="viewer-menu",
            password="secret",
            role=Usuario.Role.VISUALIZADOR,
            empresa_cliente=empresa,
        )
        self.client.force_login(viewer)

        response = self.client.get(reverse("ui:dashboard"))

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "Dashboard")
        self.assertContains(response, "Fila")
        self.assertContains(response, "Recebimento")
        self.assertContains(response, "Planilhas")
        self.assertContains(response, "Relatorios")
        self.assertNotContains(response, "Diagnostico")
        self.assertNotContains(response, "Configuracoes")
        self.assertNotContains(response, 'href="/admin/"')

    def test_sobre_renders_product_scope(self):
        response = self.client.get(reverse("ui:sobre"))

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "Preview A10")
        self.assertContains(response, "about-summary")
        self.assertContains(response, "Escopo do MVP")
        self.assertContains(response, "Garantias preservadas")
        self.assertContains(response, "Fora do MVP")
        self.assertContains(response, "XML fiscal imutavel")
        self.assertNotContains(response, "Preview A7")
