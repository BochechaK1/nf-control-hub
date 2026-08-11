from pathlib import Path
from uuid import uuid4
import hashlib
import re

from django.conf import settings
from django.db import transaction
from django.utils import timezone
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill

from audit.models import EventoAuditoria
from audit.services import registrar_evento
from files.models import Arquivo

from .models import Exportacao


COPIA_PREENCHIDA_OBSERVACAO = "Copia preenchida NFCH v2 com dados de faturamento por item."
FATURAMENTO_HEADERS = [
    "NFCH - Qtd faturada",
    "NFCH - Numero NF",
    "NFCH - Serie NF",
    "NFCH - Data faturamento",
    "NFCH - Valor total NF",
]


def readable_export_filename(associacao) -> str:
    original_name = Path(associacao.pedido.arquivo.nome_original).stem
    loja = associacao.pedido.loja.codigo or associacao.pedido.loja.nome
    nota = associacao.nota_fiscal.numero
    serie = associacao.nota_fiscal.serie
    raw_name = f"{original_name}_NF_{nota}_SERIE_{serie}_{loja}_preenchida.xlsx"
    ascii_name = raw_name.encode("ascii", "ignore").decode("ascii")
    cleaned = re.sub(r"[^A-Za-z0-9._-]+", "_", ascii_name).strip("._")
    return cleaned[:180] or f"copia_preenchida_nf_{nota}.xlsx"


def export_relative_path(filename: str) -> Path:
    now = timezone.localtime()
    return Path("exportacao") / f"{now:%Y}" / f"{now:%m}" / filename


def _load_order_workbook(associacao) -> Workbook:
    original_path = settings.NFCH_STORAGE_ROOT / associacao.pedido.arquivo.caminho_relativo
    if original_path.exists():
        return load_workbook(filename=original_path)

    workbook = Workbook()
    ws = workbook.active
    ws.title = "Pedido"
    ws.append(["Codigo", "Produto", "Quantidade"])
    for item in associacao.pedido.itens.order_by("aba", "linha", "id"):
        ws.append([item.codigo_original, item.produto_original, item.quantidade_original])
    return workbook


def _prepare_nfch_columns(ws, header_row: int) -> int:
    start_col = ws.max_column + 2
    header_fill = PatternFill("solid", fgColor="D9EAF7")
    header_font = Font(bold=True, color="0F172A")
    for offset, header in enumerate(FATURAMENTO_HEADERS):
        cell = ws.cell(row=header_row, column=start_col + offset, value=header)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center")
        ws.column_dimensions[cell.column_letter].width = max(len(header) + 2, 18)
    return start_col


def _fill_billing_columns(*, workbook: Workbook, associacao) -> None:
    nota = associacao.nota_fiscal
    data_faturamento = timezone.localtime(nota.data_emissao).date()
    allocations = (
        associacao.alocacoes.filter(status="VIGENTE")
        .select_related("item_pedido", "item_nota_fiscal")
        .order_by("item_pedido__aba", "item_pedido__linha", "id")
    )
    allocations_by_sheet: dict[str, list] = {}
    for allocation in allocations:
        allocations_by_sheet.setdefault(allocation.item_pedido.aba, []).append(allocation)

    for sheet_name, sheet_allocations in allocations_by_sheet.items():
        if sheet_name not in workbook.sheetnames:
            continue
        ws = workbook[sheet_name]
        first_item_row = min(allocation.item_pedido.linha for allocation in sheet_allocations)
        header_row = max(1, first_item_row - 1)
        start_col = _prepare_nfch_columns(ws, header_row)

        for allocation in sheet_allocations:
            row = allocation.item_pedido.linha
            values = [
                allocation.quantidade,
                nota.numero,
                nota.serie,
                data_faturamento,
                nota.valor_total,
            ]
            for offset, value in enumerate(values):
                cell = ws.cell(row=row, column=start_col + offset, value=value)
                if offset == 0:
                    cell.number_format = "0.0000"
                elif offset == 3:
                    cell.number_format = "DD/MM/YYYY"
                elif offset == 4:
                    cell.number_format = '#,##0.00'


@transaction.atomic
def gerar_copia_preenchida_associacao(*, associacao, usuario) -> Exportacao:
    existing = (
        associacao.exportacoes.filter(
            tipo=Exportacao.Tipo.COPIA_PREENCHIDA,
            observacao=COPIA_PREENCHIDA_OBSERVACAO,
        )
        .select_related("arquivo")
        .first()
    )
    if existing:
        return existing

    workbook = _load_order_workbook(associacao)
    _fill_billing_columns(workbook=workbook, associacao=associacao)

    display_filename = readable_export_filename(associacao)
    storage_filename = f"{uuid4().hex}_{display_filename}"
    relative_path = export_relative_path(storage_filename)
    absolute_path = settings.NFCH_STORAGE_ROOT / relative_path
    absolute_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(absolute_path)
    workbook.close()

    content = absolute_path.read_bytes()
    digest = hashlib.sha256(content).hexdigest()
    arquivo = Arquivo.objects.create(
        empresa_cliente=associacao.empresa_cliente,
        tipo=Arquivo.Tipo.EXPORTACAO,
        nome_original=display_filename,
        mime_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        tamanho=len(content),
        hash_sha256=digest,
        caminho_relativo=relative_path.as_posix(),
        usuario_importacao=usuario,
        politica_retencao="COPIA_PREENCHIDA_MVP",
    )
    exportacao = Exportacao.objects.create(
        empresa_cliente=associacao.empresa_cliente,
        associacao=associacao,
        arquivo=arquivo,
        tipo=Exportacao.Tipo.COPIA_PREENCHIDA,
        gerada_por=usuario,
        observacao=COPIA_PREENCHIDA_OBSERVACAO,
    )
    registrar_evento(
        empresa_cliente=associacao.empresa_cliente,
        usuario=usuario,
        origem=EventoAuditoria.Origem.USUARIO,
        entidade=exportacao,
        acao="EXPORTACAO_GERADA",
        depois={
            "associacao_id": associacao.id,
            "arquivo_id": arquivo.id,
            "hash_sha256": arquivo.hash_sha256,
        },
    )
    return exportacao
