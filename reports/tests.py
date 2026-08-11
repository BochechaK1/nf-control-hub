from decimal import Decimal
from pathlib import Path
from tempfile import TemporaryDirectory

from django.test import Client, TestCase, override_settings
from django.urls import reverse
from django.utils import timezone
from openpyxl import Workbook, load_workbook

from accounts.models import Usuario, UsuarioLoja
from approvals.services import aprovar_candidata
from files.models import Arquivo
from invoices.models import ItemNotaFiscal, NotaFiscal
from matching.services import recalcular_conferencia_para_nota
from orders.models import ItemPedido, Pedido
from organizations.models import EmpresaCliente, Fornecedor, Loja

from .models import Exportacao


class ExportAndDownloadTests(TestCase):
    def setUp(self):
        self.tmp = TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        self.override = override_settings(NFCH_STORAGE_ROOT=Path(self.tmp.name))
        self.override.enable()
        self.addCleanup(self.override.disable)

        self.empresa = EmpresaCliente.objects.create(nome="Destiny In", codigo="destiny")
        self.loja = Loja.objects.create(
            empresa_cliente=self.empresa,
            codigo="matriz",
            nome="Matriz",
            cnpj="11.111.111/0001-11",
        )
        self.outra_loja = Loja.objects.create(
            empresa_cliente=self.empresa,
            codigo="filial",
            nome="Filial",
            cnpj="22.111.111/0001-11",
        )
        self.fornecedor = Fornecedor.objects.create(
            empresa_cliente=self.empresa,
            nome="Coral",
            cnpj="22.222.222/0001-22",
        )
        self.admin = Usuario.objects.create_superuser(username="kauan", password="secret")
        self.viewer = Usuario.objects.create_user(
            username="viewer",
            password="secret",
            role=Usuario.Role.VISUALIZADOR,
            empresa_cliente=self.empresa,
        )
        UsuarioLoja.objects.create(usuario=self.viewer, loja=self.loja)
        self.file_counter = 0

    def arquivo(self, tipo):
        self.file_counter += 1
        return Arquivo.objects.create(
            empresa_cliente=self.empresa,
            tipo=tipo,
            nome_original=f"arquivo-{self.file_counter}.xlsx",
            tamanho=10,
            hash_sha256=f"{self.file_counter:064d}",
            caminho_relativo=f"arquivo-{self.file_counter}.xlsx",
        )

    def make_approved_association(self):
        arquivo_pedido = self.arquivo(Arquivo.Tipo.PLANILHA_PEDIDO)
        workbook = Workbook()
        ws = workbook.active
        ws.title = "Pedido"
        ws.append(["Codigo", "Produto", "Quantidade"])
        ws.append(["000123", "Tinta branco", 10])
        original_path = Path(self.tmp.name) / arquivo_pedido.caminho_relativo
        workbook.save(original_path)
        workbook.close()

        pedido = Pedido.objects.create(
            empresa_cliente=self.empresa,
            loja=self.loja,
            fornecedor=self.fornecedor,
            arquivo=arquivo_pedido,
            numero_pedido="P-001",
            data_operacional=timezone.localdate(),
        )
        ItemPedido.objects.create(
            pedido=pedido,
            empresa_cliente=self.empresa,
            fornecedor=self.fornecedor,
            aba="Pedido",
            linha=2,
            codigo_original="000123",
            codigo_normalizado="123",
            produto_original="Tinta branco",
            produto_normalizado="TINTA BRANCO",
            quantidade_original="10",
            quantidade_pedida=Decimal("10"),
            unidade_original="UN",
            unidade_normalizada="UN",
            saldo_cache=Decimal("10"),
        )
        nota = NotaFiscal.objects.create(
            empresa_cliente=self.empresa,
            arquivo_xml=self.arquivo(Arquivo.Tipo.XML_NFE),
            fornecedor=self.fornecedor,
            loja=self.loja,
            chave_acesso="1".zfill(44),
            modelo="55",
            numero="100",
            serie="1",
            data_emissao=timezone.now(),
            emitente_cnpj=self.fornecedor.cnpj_normalizado,
            emitente_nome=self.fornecedor.nome,
            destinatario_cnpj=self.loja.cnpj_normalizado,
            destinatario_nome=self.loja.nome,
            valor_total=Decimal("999.00"),
        )
        ItemNotaFiscal.objects.create(
            nota_fiscal=nota,
            empresa_cliente=self.empresa,
            numero_item=1,
            codigo_original="000123",
            codigo_normalizado="123",
            descricao_original="Tinta branco",
            descricao_normalizada="TINTA BRANCO",
            unidade_original="UN",
            unidade_normalizada="UN",
            quantidade_original="10",
            quantidade=Decimal("10"),
        )
        candidate = recalcular_conferencia_para_nota(nota).candidatas.first()
        return aprovar_candidata(candidata=candidate, usuario=self.admin)

    def test_approval_generates_filled_copy_export(self):
        association = self.make_approved_association()

        export = Exportacao.objects.get(associacao=association)
        path = Path(self.tmp.name) / export.arquivo.caminho_relativo
        self.assertEqual(export.tipo, Exportacao.Tipo.COPIA_PREENCHIDA)
        self.assertEqual(export.arquivo.nome_original, "arquivo-1_NF_100_SERIE_1_matriz_preenchida.xlsx")
        self.assertTrue(Path(export.arquivo.caminho_relativo).name.endswith(export.arquivo.nome_original))
        self.assertTrue(path.exists())
        self.assertGreater(export.arquivo.tamanho, 0)
        workbook = load_workbook(path)
        try:
            ws = workbook["Pedido"]
            self.assertEqual(ws["E1"].value, "NFCH - Qtd faturada")
            self.assertEqual(ws["F1"].value, "NFCH - Numero NF")
            self.assertEqual(ws["H1"].value, "NFCH - Data faturamento")
            self.assertEqual(ws["I1"].value, "NFCH - Valor total NF")
            self.assertEqual(ws["E2"].value, 10)
            self.assertEqual(ws["F2"].value, "100")
            self.assertEqual(ws["G2"].value, "1")
            self.assertEqual(ws["H2"].value.date(), timezone.localtime(association.nota_fiscal.data_emissao).date())
            self.assertEqual(ws["I2"].value, 999)
        finally:
            workbook.close()

    def test_authorized_viewer_can_download_store_file(self):
        association = self.make_approved_association()
        export = Exportacao.objects.get(associacao=association)
        client = Client()
        client.force_login(self.viewer)

        response = client.get(reverse("ui:download_arquivo", args=[export.arquivo.id]))

        try:
            self.assertEqual(response.status_code, 200)
            self.assertEqual(response["Content-Type"], "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
        finally:
            response.close()

    def test_reports_export_list_identifies_original_spreadsheet(self):
        self.make_approved_association()
        client = Client()
        client.force_login(self.admin)

        response = client.get(reverse("ui:relatorios"))

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "arquivo-1.xlsx")
        self.assertContains(response, "NF 100/1")
        self.assertContains(response, "arquivo-1_NF_100_SERIE_1_matriz_preenchida.xlsx")

    def test_viewer_cannot_download_unlinked_store_file(self):
        association = self.make_approved_association()
        association.pedido.loja = self.outra_loja
        association.pedido.save(update_fields=["loja"])
        export = Exportacao.objects.get(associacao=association)
        client = Client()
        client.force_login(self.viewer)

        response = client.get(reverse("ui:download_arquivo", args=[export.arquivo.id]))

        self.assertEqual(response.status_code, 404)
